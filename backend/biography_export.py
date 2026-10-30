"""
biography_export.py PDF and DOCX export for Thought Biography
"""
import io
import re
from datetime import datetime


def export_pdf(full_content: str, display_name: str, timestamp: str) -> bytes:
    """Generate PDF using WeasyPrint with system-safe fonts (no Google Fonts CDN)."""
    from weasyprint import HTML, CSS

    html_body = _md_to_html(full_content)

    html = f"""<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  body {{
    font-family: Georgia, "Times New Roman", Times, serif;
    font-size: 12.5pt;
    line-height: 1.85;
    color: #1a1a1a;
    max-width: 640px;
    margin: 0 auto;
    padding: 56px 0;
  }}
  h1 {{
    font-size: 20pt;
    font-weight: normal;
    font-style: italic;
    margin: 0 0 6px 0;
    letter-spacing: 0.01em;
  }}
  h2 {{
    font-size: 13pt;
    font-weight: normal;
    font-style: italic;
    color: #444;
    border-bottom: 1px solid #ccc;
    padding-bottom: 4px;
    margin: 40px 0 16px 0;
  }}
  p {{
    margin: 0 0 13px 0;
    text-align: justify;
    hyphens: auto;
  }}
  hr {{
    border: none;
    border-top: 1px solid #ccc;
    margin: 36px 0;
  }}
  .meta {{
    color: #888;
    font-size: 10.5pt;
    font-style: italic;
    margin-bottom: 48px;
    display: block;
  }}
  strong {{ font-weight: 600; }}
  em {{ font-style: italic; }}
  @page {{
    margin: 2.2cm 2.5cm;
  }}
</style>
</head>
<body>
{html_body}
</body>
</html>"""

    return HTML(string=html).write_pdf()


def export_docx(full_content: str, display_name: str, timestamp: str) -> bytes:
    """Generate a well-formatted DOCX using python-docx."""
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin    = Inches(1.25)
        section.bottom_margin = Inches(1.25)
        section.left_margin   = Inches(1.4)
        section.right_margin  = Inches(1.4)

    # Default paragraph style
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(8)

    def add_title(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Georgia"
        run.font.size = Pt(22)
        run.font.italic = True
        run.font.bold = False
        p.paragraph_format.space_after = Pt(4)
        return p

    def add_meta(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Georgia"
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        run.font.italic = True
        p.paragraph_format.space_after = Pt(36)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.font.name = "Georgia"
        run.font.size = Pt(13)
        run.font.italic = True
        run.font.bold = False
        p.paragraph_format.space_before = Pt(28)
        p.paragraph_format.space_after = Pt(10)
        # Bottom border
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "4")
        bottom.set(qn("w:space"), "1")
        bottom.set(qn("w:color"), "CCCCCC")
        pBdr.append(bottom)
        pPr.append(pBdr)
        return p

    def add_hr():
        p = doc.add_paragraph()
        run = p.add_run("─" * 28)
        run.font.name = "Georgia"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(20)
        return p

    def add_body(text):
        """Add a paragraph with inline bold/italic support."""
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        # Parse inline formatting: **bold**, *italic*, ***bold-italic***
        tokens = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*)', text)
        for token in tokens:
            if token.startswith("***") and token.endswith("***"):
                run = p.add_run(token[3:-3])
                run.bold = True
                run.italic = True
            elif token.startswith("**") and token.endswith("**"):
                run = p.add_run(token[2:-2])
                run.bold = True
            elif token.startswith("*") and token.endswith("*"):
                run = p.add_run(token[1:-1])
                run.italic = True
            else:
                run = p.add_run(token)
            run.font.name = "Georgia"
            run.font.size = Pt(12)
        return p

    # --- Parse and render ---
    lines = full_content.split("\n")
    i = 0
    first_h1_done = False

    while i < len(lines):
        line = lines[i]

        if line.startswith("# "):
            if not first_h1_done:
                add_title(line[2:].strip())
                first_h1_done = True
            i += 1

        elif line.startswith("## "):
            add_h2(line[3:].strip())
            i += 1

        elif line.strip() == "---":
            add_hr()
            i += 1

        elif line.strip() == "":
            i += 1

        else:
            add_body(line.strip())
            i += 1

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


def _md_to_html(md: str) -> str:
    """Convert markdown to HTML for PDF export — handles inline bold/italic."""
    lines = []
    for line in md.split("\n"):
        stripped = line.strip()
        if stripped.startswith("# "):
            lines.append(f"<h1>{_inline_md(stripped[2:])}</h1>")
        elif stripped.startswith("## "):
            lines.append(f"<h2>{_inline_md(stripped[3:])}</h2>")
        elif stripped == "---":
            lines.append("<hr>")
        elif stripped == "":
            lines.append("")
        else:
            lines.append(f"<p>{_inline_md(stripped)}</p>")
    return "\n".join(lines)


def _inline_md(text: str) -> str:
    """Convert inline **bold**, *italic*, ***bold-italic*** to HTML."""
    text = re.sub(r'\*\*\*(.*?)\*\*\*', r'<strong><em>\1</em></strong>', text)
    text = re.sub(r'\*\*(.*?)\*\*', r'<strong>\1</strong>', text)
    text = re.sub(r'\*(.*?)\*', r'<em>\1</em>', text)
    return text
